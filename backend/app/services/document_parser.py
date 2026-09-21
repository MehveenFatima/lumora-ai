import io
import re
from typing import List, Dict, Any, Optional
from backend.app.core.logging import get_logger

logger = get_logger(__name__)


class ParsedPage:
    def __init__(
        self,
        page_number: int,
        text: str,
        heading: Optional[str] = None,
        section_breadcrumbs: Optional[List[str]] = None,
        has_table: bool = False,
        tables_markdown: Optional[List[str]] = None,
        has_diagram: bool = False,
        figure_captions: Optional[List[str]] = None,
    ):
        self.page_number = page_number
        self.text = text.strip()
        self.heading = heading
        self.section_breadcrumbs = section_breadcrumbs or []
        self.has_table = has_table
        self.tables_markdown = tables_markdown or []
        self.has_diagram = has_diagram
        self.figure_captions = figure_captions or []


class DocumentParserService:
    def parse(self, file_bytes: bytes, file_type: str) -> List[ParsedPage]:
        f_type = file_type.lower().strip(".")
        if f_type == "pdf":
            return self._parse_pdf(file_bytes)
        elif f_type == "docx":
            return self._parse_docx(file_bytes)
        elif f_type == "pptx":
            return self._parse_pptx(file_bytes)
        elif f_type == "txt":
            return self._parse_txt(file_bytes)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    def render_page_image(self, file_bytes: bytes, page_number: int, dpi: int = 150) -> Optional[bytes]:
        """Renders a specific PDF page to PNG bytes for vision grounding."""
        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            if 1 <= page_number <= len(doc):
                pix = doc[page_number - 1].get_pixmap(dpi=dpi)
                png_bytes = pix.tobytes("png")
                doc.close()
                return png_bytes
            doc.close()
        except Exception as e:
            logger.warning(f"Could not render page {page_number} to image: {e}")
        return None

    def _parse_pdf(self, file_bytes: bytes) -> List[ParsedPage]:
        pages: List[ParsedPage] = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                page_num = page_idx + 1
                raw_text = (page.get_text("text") or "").replace("\x00", "")
                
                # 1. Table Extraction using PyMuPDF find_tables()
                tables_md: List[str] = []
                try:
                    tab_finder = page.find_tables()
                    if tab_finder and tab_finder.tables:
                        for tab in tab_finder.tables:
                            extracted = tab.extract()
                            if extracted and len(extracted) >= 2:
                                headers = [str(c or "").strip() for c in extracted[0]]
                                md_lines = ["| " + " | ".join(headers) + " |"]
                                md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                                for row in extracted[1:]:
                                    row_vals = [str(c or "").strip().replace("\n", " ") for c in row]
                                    md_lines.append("| " + " | ".join(row_vals) + " |")
                                tables_md.append("\n".join(md_lines))
                except Exception as te:
                    logger.debug(f"find_tables warning on page {page_num}: {te}")

                # 2. Heading Extraction & Caption Detection
                lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
                heading: Optional[str] = None
                figure_captions: List[str] = []
                
                for idx, line in enumerate(lines[:5]):
                    # If line looks like a title or section heading
                    if (
                        len(line) >= 3
                        and len(line) < 90
                        and not line.isdigit()
                        and not line.lower().startswith("page ")
                        and not line.lower().startswith("chapter ") and len(line) > 5
                    ):
                        heading = line
                        break
                    elif (
                        line.isupper()
                        and len(line) < 70
                        and not line.isdigit()
                    ):
                        heading = line
                        break

                for line in lines:
                    if re.match(r"^(figure|fig\.|diagram|table|chart|graph)\s+\d+[\:\.\-]", line, re.IGNORECASE):
                        figure_captions.append(line)

                # 3. Image / Diagram Detection
                has_images = len(page.get_images()) > 0 or len(figure_captions) > 0
                has_table = len(tables_md) > 0 or any(re.match(r"^table\s+\d+", l, re.IGNORECASE) for l in lines)

                # Combine text and extracted table representations
                combined_text = raw_text
                if tables_md:
                    combined_text += "\n\n### Extracted Tables:\n" + "\n\n".join(tables_md)

                if combined_text.strip():
                    pages.append(
                        ParsedPage(
                            page_number=page_num,
                            text=combined_text,
                            heading=heading,
                            section_breadcrumbs=[heading] if heading else [],
                            has_table=has_table,
                            tables_markdown=tables_md,
                            has_diagram=has_images,
                            figure_captions=figure_captions,
                        )
                    )
            doc.close()
        except ImportError:
            logger.warning("PyMuPDF (fitz) not available. Falling back to plain text extraction.")
            text = file_bytes.decode("utf-8", errors="ignore")
            pages.append(ParsedPage(page_number=1, text=text))
        except Exception as e:
            logger.error(f"Failed to parse PDF file: {e}")
            raise ValueError(f"Failed to extract text from PDF: {e}")

        if not pages:
            pages.append(ParsedPage(page_number=1, text="[Empty Document]"))
        return pages

    def _parse_docx(self, file_bytes: bytes) -> List[ParsedPage]:
        pages: List[ParsedPage] = []
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            
            # Process paragraphs with heading breadcrumbs
            blocks = []
            tables_md = []
            current_heading: Optional[str] = None
            breadcrumbs: List[str] = []

            for para in doc.paragraphs:
                p_text = para.text.strip()
                if not p_text:
                    continue
                style_name = (para.style.name if para.style else "").lower()
                if "heading 1" in style_name or "title" in style_name:
                    current_heading = p_text
                    breadcrumbs = [p_text]
                elif "heading 2" in style_name:
                    if breadcrumbs:
                        breadcrumbs = [breadcrumbs[0], p_text]
                    else:
                        breadcrumbs = [p_text]
                    current_heading = p_text
                elif len(p_text) < 70 and (p_text.isupper() or p_text.startswith("#")):
                    current_heading = p_text.lstrip("#").strip()

                blocks.append(p_text)

            # Process Tables
            for table in doc.tables:
                if not table.rows:
                    continue
                headers = [c.text.strip().replace("\n", " ") for c in table.rows[0].cells]
                t_lines = ["| " + " | ".join(headers) + " |"]
                t_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for row in table.rows[1:]:
                    vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                    t_lines.append("| " + " | ".join(vals) + " |")
                t_md = "\n".join(t_lines)
                tables_md.append(t_md)
                blocks.append(t_md)

            # Partition into bounded ~400-word logical pages
            full_text = "\n\n".join(blocks)
            words = full_text.split()
            page_size = 400
            for i in range(0, max(len(words), 1), page_size):
                page_words = words[i : i + page_size]
                page_text = " ".join(page_words)
                page_num = (i // page_size) + 1
                if page_text.strip():
                    pages.append(
                        ParsedPage(
                            page_number=page_num,
                            text=page_text,
                            heading=current_heading,
                            section_breadcrumbs=breadcrumbs,
                            has_table=len(tables_md) > 0,
                            tables_markdown=tables_md,
                        )
                    )
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise ValueError(f"Failed to parse DOCX document: {e}")

        if not pages:
            pages.append(ParsedPage(page_number=1, text="[Empty Document]"))
        return pages

    def _parse_pptx(self, file_bytes: bytes) -> List[ParsedPage]:
        pages: List[ParsedPage] = []
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            for idx, slide in enumerate(prs.slides):
                slide_num = idx + 1
                slide_texts = []
                heading = None
                tables_md = []
                has_diagram = False

                # 1. Slide Title
                if slide.shapes.title and slide.shapes.title.text:
                    heading = slide.shapes.title.text.strip()
                    slide_texts.append(f"# {heading}")

                # 2. Slide Shapes & Text Frames
                for shape in slide.shapes:
                    # Diagram / picture detection
                    if shape.shape_type in (13, 6):  # 13 = PICTURE, 6 = GROUP
                        has_diagram = True

                    # Table in slide
                    if shape.has_table:
                        tbl = shape.table
                        if tbl.rows:
                            h_cells = [c.text.strip().replace("\n", " ") for c in tbl.rows[0].cells]
                            t_lines = ["| " + " | ".join(h_cells) + " |"]
                            t_lines.append("| " + " | ".join(["---"] * len(h_cells)) + " |")
                            for r in tbl.rows[1:]:
                                v_cells = [c.text.strip().replace("\n", " ") for c in r.cells]
                                t_lines.append("| " + " | ".join(v_cells) + " |")
                            t_md = "\n".join(t_lines)
                            tables_md.append(t_md)
                            slide_texts.append(t_md)

                    # Text Frame
                    if shape.has_text_frame:
                        if shape == slide.shapes.title:
                            continue
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                if not heading:
                                    heading = t[:80]
                                slide_texts.append(t)

                text = "\n\n".join(slide_texts)
                if text.strip():
                    pages.append(
                        ParsedPage(
                            page_number=slide_num,
                            text=text,
                            heading=heading,
                            section_breadcrumbs=[heading] if heading else [],
                            has_table=len(tables_md) > 0,
                            tables_markdown=tables_md,
                            has_diagram=has_diagram,
                        )
                    )
        except Exception as e:
            logger.error(f"Failed to parse PPTX: {e}")
            raise ValueError(f"Failed to parse presentation: {e}")

        if not pages:
            pages.append(ParsedPage(page_number=1, text="[Empty Document]"))
        return pages

    def _parse_txt(self, file_bytes: bytes) -> List[ParsedPage]:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")

        page_size = 2000
        pages: List[ParsedPage] = []
        for i in range(0, max(len(text), 1), page_size):
            chunk = text[i : i + page_size]
            if chunk.strip():
                pages.append(ParsedPage(page_number=(i // page_size) + 1, text=chunk))

        if not pages:
            pages.append(ParsedPage(page_number=1, text="[Empty Document]"))
        return pages


document_parser = DocumentParserService()

